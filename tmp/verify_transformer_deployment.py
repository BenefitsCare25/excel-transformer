"""Verify downloaded listing codes using synthetic clinic data only."""
import argparse
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import threading

import pandas as pd
from openpyxl import load_workbook
import requests


def verify(base_url):
    session = requests.Session()
    health = session.get(base_url + '/health', timeout=60)
    health.raise_for_status()
    source = pd.DataFrame({
        'S/N': [1, 2, 3, 4, 5],
        'Clinic ID': ['P100', None, 'P100', 'P100-1', 'DROP'],
        'Clinic Name': ['Synthetic Alpha', 'Synthetic Beta', 'Synthetic Gamma', 'Synthetic Delta', 'Synthetic Removed'],
        'Region': ['NORTH'] * 5,
        'Area': ['TEST AREA'] * 5,
        'Address': ['1 TEST ROAD SINGAPORE 123456'] * 5,
        'Postal Code': ['123456'] * 5,
        'Tel No.': ['61234567'] * 5,
    })
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        for sheet_name in ['GP Panel', 'Dental Panel', 'TCM Panel', 'SP Clinic']:
            panel = source.copy()
            if sheet_name == 'GP Panel':
                panel.loc[[1, 3], 'Address'] = '2 JALAN TEST 80000 JOHOR MALAYSIA'
                panel.loc[[1, 3], 'Postal Code'] = '80000'
            panel.to_excel(writer, sheet_name=sheet_name, index=False)
        pd.DataFrame({
            'S/N': [1], 'Region': ['NORTH'], 'Area': ['TEST AREA'],
            'Provider Code': ['DROP'], 'Clinic Name': ['Synthetic Removed'], 'Postal Code': ['123456'],
            'Tel No.': ['61234567'],
        }).to_excel(writer, sheet_name='Terminated Clinics', index=False)
    response = session.post(base_url + '/upload', files={
        'file': ('synthetic-code-check.xlsx', stream.getvalue()),
    }, data={'use_google_api': 'false'}, timeout=120)
    response.raise_for_status()
    result = response.json()
    assert result['total_records'] == 16, result
    assert result['terminated_clinics_filtered'] == 4, result
    assert len(result['download_urls']) == 5, result
    for output, url in zip(result['results'], result['download_urls']):
        download = session.get(base_url + url, timeout=60)
        download.raise_for_status()
        workbook = load_workbook(BytesIO(download.content), data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.values)
            assert rows[0][0] == 'Code', rows[0]
            codes = [row[0] for row in rows[1:]]
            assert codes == list(range(1, output['records_processed'] + 1)), codes
            assert all(sheet.cell(row, 1).data_type == 'n' for row in range(2, sheet.max_row + 1))
            assert all(row[1] and row[1] != 'Synthetic Removed' for row in rows[1:])
            print(f"PASS {output['output_filename']}: {len(codes)} complete, unique numeric codes", flush=True)
        finally:
            workbook.close()
    print('PASS source-code termination matching, every listing type and country splits', flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base-url', default='https://excel-transformer-rg.azurewebsites.net')
    parser.add_argument('--local', action='store_true', help='Run against the current backend with temporary upload/output directories.')
    args = parser.parse_args()
    if args.local:
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
        import app as backend
        from werkzeug.serving import make_server

        with tempfile.TemporaryDirectory() as directory:
            backend.UPLOAD_FOLDER = str(Path(directory) / 'uploads')
            backend.PROCESSED_FOLDER = str(Path(directory) / 'processed')
            Path(backend.UPLOAD_FOLDER).mkdir()
            Path(backend.PROCESSED_FOLDER).mkdir()
            server = make_server('127.0.0.1', 0, backend.app, threaded=True)
            threading.Thread(target=server.serve_forever, daemon=True).start()
            try:
                verify(f'http://127.0.0.1:{server.server_port}')
            finally:
                server.shutdown()
                server.server_close()
    else:
        verify(args.base_url.rstrip('/'))
