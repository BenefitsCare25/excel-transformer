"""RHB Excel layouts, including employee subtotals and all 14 benefit columns."""

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import (
    CLAIM_COLUMNS, CLAIM_RULES, CODE_DESCRIPTIONS, DETAIL_FIELDS, DETAIL_HEADERS,
    DETAIL_WIDTHS, PALETTE, SUMMARY_WIDTHS, UTILIZATION_FIELDS, UTILIZATION_HEADERS,
    UTILIZATION_WIDTHS,
)

MONEY_FORMAT = '#,##0.00'


def _put(sheet, row, column, value, bold=False):
    if pd.isna(value):
        value = None
    elif isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Calibri", size=11, bold=bold, color=PALETTE["text"])
    cell.alignment = Alignment(vertical="center")
    return cell


def _configure(sheet, widths, header_row=3):
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f"C{header_row + 1}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_view.showGridLines = False


def _heading(sheet, headers, row=3):
    edge = Side(style="thin", color=PALETTE["grid"])
    for column, label in enumerate(headers, 1):
        cell = _put(sheet, row, column, label, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
    sheet.row_dimensions[row].height = 32


def _row(sheet, row_number, values, money=(), dates=(), bold=False, wrap=()):
    for column, value in enumerate(values, 1):
        cell = _put(sheet, row_number, column, value, bold=bold)
        if column in money:
            cell.number_format = MONEY_FORMAT
        if column in dates:
            cell.number_format = "dd/mm/yyyy"
        if column in wrap:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    if wrap:
        # Excel does not reliably autofit wrapped rows in generated workbooks.
        lines = max((len(str(values[col - 1])) // max(1, int(sheet.column_dimensions[
            get_column_letter(col)].width) - 4) + 1 for col in wrap if pd.notna(values[col - 1])), default=1)
        sheet.row_dimensions[row_number].height = max(16, lines * 15)


def _save(workbook, outdir, filename):
    path = Path(outdir) / filename
    workbook.save(path)
    workbook.close()
    return str(path)


def write_details(claims, month, outdir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Claims Detail"
    _configure(sheet, DETAIL_WIDTHS)
    _put(sheet, 1, 1, "Claims Report", bold=True)
    _heading(sheet, DETAIL_HEADERS)
    row_number = 4
    for (_, _, name), employee in claims.groupby(["Entity", "Staff ID", "Employee Name"], sort=False):
        start = row_number
        for _, claim in employee.iterrows():
            _row(sheet, row_number, [claim[field] for field in DETAIL_FIELDS], money=(9, 10),
                 dates=(7,), wrap=(2, 3, 6, 8, 11))
            sheet.row_dimensions[row_number].outlineLevel = 1
            row_number += 1
        _put(sheet, row_number, 2, f"{name} Total", bold=True)
        for column in (9, 10):
            letter = get_column_letter(column)
            cell = sheet.cell(row_number, column, f"=SUBTOTAL(9,{letter}{start}:{letter}{row_number - 1})")
            cell.font = Font(name="Calibri", size=11, bold=True)
            cell.number_format = MONEY_FORMAT
        row_number += 1
    _put(sheet, row_number, 2, "Grand Total", bold=True)
    for column in (9, 10):
        letter = get_column_letter(column)
        cell = sheet.cell(row_number, column, f"=SUBTOTAL(9,{letter}4:{letter}{row_number - 2})")
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.number_format = MONEY_FORMAT
    return _save(workbook, outdir, f"Claims Detail Report - {month:%B %Y}.xlsx")


def _summary_headings(sheet):
    _configure(sheet, SUMMARY_WIDTHS, header_row=4)
    _put(sheet, 1, 1, "Claim Summary Report", bold=True)
    for column, label in enumerate(("Employee ID", "Employee Name", "No. of Receipt", "CCY"), 1):
        sheet.merge_cells(start_row=3, end_row=4, start_column=column, end_column=column)
        _put(sheet, 3, column, label, bold=True)
    for start, end, code in ((5, 8, "RE0235"), (9, 9, "RE0238"), (10, 18, "RE0237")):
        if start != end:
            sheet.merge_cells(start_row=3, end_row=3, start_column=start, end_column=end)
        _put(sheet, 3, start, CODE_DESCRIPTIONS[code], bold=True)
    for column, label in enumerate(CLAIM_RULES, 5):
        _put(sheet, 4, column, label, bold=True)
    for column, label in enumerate(("Grand Total", *[f"Total for {label}" for label in CODE_DESCRIPTIONS.values()]), 19):
        sheet.merge_cells(start_row=3, end_row=4, start_column=column, end_column=column)
        _put(sheet, 3, column, label, bold=True)
    edge = Side(style="thin", color=PALETTE["grid"])
    for row in sheet.iter_rows(min_row=3, max_row=4, max_col=22):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=PALETTE["benefit" if 5 <= cell.column <= 18 else "heading"])
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
    sheet.row_dimensions[3].height = 46
    sheet.row_dimensions[4].height = 112


def write_summary(claims, month, outdir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Claim Summary"
    _summary_headings(sheet)
    groups = claims.groupby(["Staff ID", "Employee Name"], sort=True)
    for row_number, ((staff_id, name), employee) in enumerate(groups, 5):
        amounts = employee.groupby("ReportType")["PaymentCents"].sum()
        codes = employee.groupby("PayrollCode")["PaymentCents"].sum()
        values = [staff_id, name, len(employee), "SGD",
                  *[int(amounts[label]) / 100 if label in amounts else None for label in CLAIM_RULES],
                  int(employee["PaymentCents"].sum()) / 100,
                  *[int(codes.get(code, 0)) / 100 for code in CODE_DESCRIPTIONS]]
        _row(sheet, row_number, values, money=range(5, 23))
    total_row = sheet.max_row + 1
    _put(sheet, total_row, 1, "Grand Total", bold=True)
    sheet.merge_cells(start_row=total_row, end_row=total_row, start_column=1, end_column=2)
    _put(sheet, total_row, 3, len(claims), bold=True)
    _put(sheet, total_row, 4, "SGD", bold=True)
    # Static report totals remain readable by payroll importers without Excel recalculation.
    for column in range(5, 23):
        total = sum(int(round((sheet.cell(row, column).value or 0) * 100)) for row in range(5, total_row))
        _put(sheet, total_row, column, total / 100, bold=True).number_format = MONEY_FORMAT
    raw = workbook.create_sheet("Claims Data")
    _configure(raw, [max(18, min(45, len(label) + 4)) for label in CLAIM_COLUMNS], header_row=1)
    _heading(raw, CLAIM_COLUMNS, row=1)
    for row_number, (_, claim) in enumerate(claims.iterrows(), 2):
        _row(raw, row_number, [claim[column] for column in CLAIM_COLUMNS],
             money=(15, 17, 18), dates=(12, 20))
    raw.auto_filter.ref = f"A1:U{raw.max_row}"
    return _save(workbook, outdir, f"Claim Summary Report - {month:%B %Y}.xlsx")


def write_utilization(utilization, month, outdir):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = str(month.year)
    _configure(sheet, UTILIZATION_WIDTHS)
    _put(sheet, 1, 1, "Utilisation and Balance Report", bold=True)
    _heading(sheet, UTILIZATION_HEADERS)
    ordered = utilization.sort_values("Employee Name", key=lambda names: names.str.casefold(), kind="stable")
    for row_number, (_, record) in enumerate(ordered.iterrows(), 4):
        _row(sheet, row_number, [record[field] for field in UTILIZATION_FIELDS],
             money=range(7, 11), dates=(3, 4), wrap=(2, 5, 6))
    sheet.auto_filter.ref = f"A3:J{sheet.max_row}"
    return _save(workbook, outdir, f"Utilisation and Balance Report - {month:%B %Y}.xlsx")
