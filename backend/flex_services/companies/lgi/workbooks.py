"""Generate LGI Excel reports in the supplied reference layouts."""

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import (
    DETAIL_HEADERS, DETAIL_WIDTHS, SUMMARY_HEADERS, SUMMARY_WIDTHS,
    ENTITY, UTILIZATION_HEADERS, UTILIZATION_WIDTHS,
)
from .processing import policy_period

MONEY_FORMAT = "#,##0.00"
EDGE = Side(style="thin", color="000000")
GRID = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)


def put(sheet, row, column, value):
    if pd.isna(value):
        value = None
    elif isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    cell = sheet.cell(row, column, value)
    # Source strings are literal text, even when they begin with formula markers.
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def table(headers, widths, rows, money_columns, date_columns=(), header_row=1, yellow=True):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    sheet.sheet_format.defaultRowHeight = 14.4
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for col, label in enumerate(headers, 1):
        cell = put(sheet, header_row, col, label)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if yellow:
            cell.fill = PatternFill("solid", fgColor="FFFF00")
    for row, values in enumerate(rows, header_row + 1):
        for col, value in enumerate(values, 1):
            cell = put(sheet, row, col, value)
            cell.font = Font(name="Calibri", size=11)
            if col in money_columns:
                cell.number_format = MONEY_FORMAT
            if col in date_columns:
                cell.number_format = "dd/mm/yyyy"
    for row in sheet.iter_rows(min_row=header_row):
        for cell in row:
            cell.border = GRID
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"{header_row}:{header_row}"
    return workbook


def write_details(claims, month, outdir):
    rows = []
    for _, claim in claims.iterrows():
        rows.append([
            claim["Reference No."], claim["Incurred Date"], claim["Staff ID"],
            claim["Employee Name"], claim["Claimant Name"], claim["Relation"],
            claim["Claim Type"], claim["Taxable"], claim["Non Taxable"], claim["CPF"],
            claim["Converted Incurred Amt"], claim["Payment Amt"], claim["Service Provider"],
            "Paid", claim["Admin Remark"], None,
        ])
    workbook = table(DETAIL_HEADERS, DETAIL_WIDTHS, rows, (8, 9, 11, 12), (2,))
    path = Path(outdir) / f"Claims Detail _ {month:%B %Y}.xlsx"
    workbook.save(path)
    return str(path)


def write_summary(claims, month, outdir):
    # The LGI reference deliberately retains one summary row per claim.
    rows = [[c["Staff ID"], c["Employee Name"], c["Designation"], c["Cost Centre"],
             c["Department"], c["Taxable"], c["Non Taxable"], c["CPF"]]
            for _, c in claims.iterrows()]
    workbook = table(SUMMARY_HEADERS, SUMMARY_WIDTHS, rows, (6, 7))
    for cell in workbook.active[1]:
        cell.font = Font(name="Arial", size=10, bold=True)
    path = Path(outdir) / f"Claims Summary _ {month:%B %Y}.xlsx"
    workbook.save(path)
    return str(path)


def write_utilization(utilization, month, outdir):
    start, end = policy_period(month.isoformat())
    label = f"1 October {start.year}  to 30 September {end.year}"
    rows = [[label, row["Staff ID"], row["Employee Name"], row["Designation"],
             row["Cost Centre"], row["Department"], row["Total Allocation Amt"],
             row["FlexUsed"], row["TransferredFSA"], row["MonthlyClaims"],
             row["Claims Payment Amt"], None, row["Balance Available Allocation Amt"],
             row["LastDay"]] for _, row in utilization.iterrows()]
    workbook = table(UTILIZATION_HEADERS, UTILIZATION_WIDTHS, rows, range(7, 14),
                     (14,), header_row=6, yellow=False)
    sheet = workbook.active
    sheet.title = f"1 Oct {start.year}  to 30 Sep {end.year}"
    for row, title in [(1, ENTITY.upper()), (2, "Utilization Report "), (4, f"Month: {month:%B %Y}")]:
        cell = put(sheet, row, 1, title)
        cell.font = Font(name="Calibri", size=11, bold=True)
    sheet.row_dimensions[6].height = 37.2
    sheet.print_title_rows = "1:6"
    path = Path(outdir) / f"Utilization Report_{month:%B %Y}.xlsx"
    workbook.save(path)
    return str(path)
