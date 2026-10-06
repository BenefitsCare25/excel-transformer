"""Write Ceva reports in the supplied payroll and utilisation formats."""

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import (
    CODE_DESCRIPTIONS, CLAIM_RULES, FILE_COMPANY_LABEL, PALETTE,
    PAYROLL_HEADERS, PAYROLL_WIDTHS, UTILIZATION_HEADERS, UTILIZATION_WIDTHS,
)


def _table(headers, widths, rows, money_columns, currency=False, date_columns=()):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    edge = Side(style="thin", color=PALETTE["grid"])
    grid = Border(left=edge, right=edge, top=edge, bottom=edge)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_number, values in enumerate([headers, *rows], 1):
        for column, value in enumerate(values, 1):
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            cell = sheet.cell(row_number, column, value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11, bold=row_number == 1)
            cell.border = grid
            cell.alignment = Alignment(vertical="center", wrap_text=row_number == 1)
            if row_number == 1:
                cell.fill = PatternFill("solid", fgColor=PALETTE["header"])
            if column in money_columns:
                cell.number_format = '"$"#,##0.00' if currency else '#,##0.00'
            if row_number > 1 and column in date_columns:
                cell.number_format = "dd/mm/yyyy"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    return workbook


def payroll_rows(claims):
    grouped = claims.groupby(["Entity", "Staff ID", "Employee Name", "ReportType"],
                             sort=False, dropna=False)["Payment Amt"].sum().round(2)
    rows = [[entity, staff_id, name, claim_type, float(amount), CLAIM_RULES[claim_type][2],
             CODE_DESCRIPTIONS[claim_type]]
            for (entity, staff_id, name, claim_type), amount in grouped.items()]
    return sorted(rows, key=lambda row: (row[0].casefold(), row[2].casefold(),
                                         row[1], row[3].casefold()))


def write_payroll(claims, month, outdir):
    workbook = _table(PAYROLL_HEADERS, PAYROLL_WIDTHS, payroll_rows(claims), (5,))
    path = Path(outdir) / (
        f"Flexible Benefits - {month:%B %Y} reimbursement ({FILE_COMPANY_LABEL}).xlsx"
    )
    workbook.save(path)
    return str(path)


def write_utilization(utilization, month, outdir):
    ordered = utilization.sort_values("Employee Name", key=lambda values: values.str.casefold(),
                                     kind="stable")
    rows = [[row["Entity"], row["Staff ID"], row["Employee Name"], row["Date of Hire"],
             row["Wallet"], row["Total Allocation Amt"], row["Claims Payment Amt"],
             row["Pending Claims Payment Amt"], row["Balance Available Allocation Amt"]]
            for _, row in ordered.iterrows()]
    workbook = _table(UTILIZATION_HEADERS, UTILIZATION_WIDTHS, rows, range(6, 10),
                      currency=True, date_columns=(4,))
    workbook.active.title = str(month.year)
    path = Path(outdir) / (
        f"Utilisation report - {month:%B %Y} Reimbursement ({FILE_COMPANY_LABEL}).xlsx"
    )
    workbook.save(path)
    return str(path)
