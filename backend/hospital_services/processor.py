"""Process scanned hospital bills without sending raw pages to a remote service."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import gc
import math
from pathlib import Path
import re
from typing import Callable

import cv2
import fitz
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from . import checkpoints
from .fields import COMPETING_HEADERS, REF_PATTERN, extract_page
from .merging import merge_pages
from .ocr import MAX_IMAGE_SIDE, create_engine, read_lines, retry_fields
from .ocr_types import OcrItem
from .workbook_review import add_review_sheet


HEADERS = (
    "Bill Ref No.", "Bill Date", "HRN", "Visit Date",
    "Total Amount (After Govt Subsidy)", "Payable by Medishield Life",
    "Payable by Medisave", "Total Amount Payable\nCash",
)
RENDER_SCALE = 2.0
ID_PATTERN = re.compile(r"(?<![A-Z0-9])[STFGM](?:[\s.\-]*[0-9XOIL*]){7}[\s.\-]*[A-Z0-9X*](?![A-Z0-9])", re.I)
MASKED_ID_PATTERN = re.compile(r"(?<![A-Z0-9])[A-Z0-9]*[X*]{3,}[A-Z0-9]*(?![A-Z0-9])", re.I)


def _identifier(text: str) -> bool:
    if ID_PATTERN.search(text):
        return True
    compact = re.sub(r"[\s.\-]", "", text.upper())
    return any(
        len(match.group()) >= 8 and len(match.group()) <= 12
        for match in MASKED_ID_PATTERN.finditer(compact)
    )


def _identifier_regions(lines: list[OcrItem], image_width: int) -> list[tuple[float, float, float, float]]:
    regions = []
    for item in lines:
        if "NRIC" not in item.text.upper():
            continue
        label = re.match(r"\s*NRIC(?:\s*/\s*(?:FIN|MRN))*\s*[:.]?\s*", item.text, re.I)
        if label is None:
            raise ValueError("Identifier field layout could not be established for safe redaction.")
        height = item.height
        boundaries = [other.left for other in lines if other.left > item.left + height
                      and abs(other.center_y - item.center_y) <= height * 3.5
                      and any(re.sub(r"[^A-Z0-9]", "", other.text.upper()).startswith(heading)
                              for heading in COMPETING_HEADERS)]
        right = min(boundaries, default=float(image_width))
        if right <= item.right:
            raise ValueError("Identifier field boundary could not be established for safe redaction.")
        regions.append((item.left - 3, item.bottom + 1, right - 1, item.bottom + height * 2.5))
        if item.text[label.end():].strip():
            regions.append((item.left - 3, item.top - 3, item.right + 3, item.bottom + 3))
        row_ends = (item.row_y_at(item.right), item.row_y_at(right))
        regions.append((item.right - height * .6, min(row_ends) - height * 1.3,
                        right - 1, max(row_ends) + height * 1.3))
        inline = [other for other in lines if other is not item
                  and item.right - height * .6 <= other.left < right
                  and not any(re.sub(r"[^A-Z0-9]", "", other.text.upper()).startswith(heading)
                              for heading in COMPETING_HEADERS)
                  and abs(other.center_y - item.row_y_at(other.left)) <= max(height, other.height) * .8]
        if any(other.right > right for other in inline):
            raise ValueError("Identifier value crosses a field boundary; safe redaction could not be established.")
        regions.extend((other.left - 3, other.top - 3, other.right + 3, other.bottom + 3) for other in inline)
    return regions


def _redact_image(image: np.ndarray, lines: list[OcrItem]) -> int:
    regions = _identifier_regions(lines, image.shape[1])
    for item in lines:
        if not _identifier(item.text):
            continue
        regions.append((item.left - 5, item.top - 4, item.right + 6, item.bottom + 5))
    count = 0
    for left, top, right, bottom in regions:
        x1, x2 = max(0, int(left)), min(image.shape[1], int(right))
        y1, y2 = max(0, int(top)), min(image.shape[0], int(bottom))
        if x2 <= x1 or y2 <= y1:
            continue
        cv2.rectangle(image, (x1, y1), (x2, y2), (0, 0, 0), -1)
        count += 1
    return count


def _open_document(source: Path) -> fitz.Document:
    with source.open("rb") as handle:
        if handle.read(5) != b"%PDF-":
            raise ValueError("Upload a valid PDF file.")
    try:
        document = fitz.open(str(source), filetype="pdf")
    except Exception as exc:
        raise ValueError("Could not open the PDF.") from exc
    try:
        if document.is_encrypted:
            raise ValueError("PDF is encrypted. Upload an unencrypted copy.")
        if not len(document):
            raise ValueError("PDF contains no pages.")
    except Exception:
        document.close()
        raise
    return document


def _process_page(page: fitz.Page, page_number: int, engine) -> tuple[np.ndarray, int, dict | None]:
    scale = min(RENDER_SCALE, MAX_IMAGE_SIDE / max(page.rect.width, page.rect.height, 1))
    pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
        pixmap.height, pixmap.width, 3).copy()
    original_lines = read_lines(image, engine)
    try:
        count = _redact_image(image, original_lines)
    except ValueError as exc:
        raise ValueError(f"Page {page_number}: {exc}") from exc
    clean_lines = retry_fields(image, engine, read_lines(image, engine))
    if any(_identifier(reading.text) for item in clean_lines for reading in item.readings):
        raise ValueError(f"Page {page_number}: an identifier remains after redaction.")
    return image, count, extract_page(clean_lines, page_number, image)


def process_pdf(
    source: Path, target: Path, work_dir: Path, on_page: Callable[[int, int], None] | None = None
) -> tuple[list[dict], int, list[str]]:
    """OCR ``source`` page by page, writing the redacted copy to ``target``.

    Each finished page is checkpointed in ``work_dir``, so a restarted worker resumes after the
    last completed page instead of repeating the whole document.
    """
    extracted = []
    redaction_count = 0
    unreadable = []
    with _open_document(source) as document:
        engine = None
        if on_page:
            on_page(0, len(document))
        for page_number, page in enumerate(document, 1):
            saved = checkpoints.load_page(work_dir, page_number)
            if saved is None:
                engine = engine or create_engine()
                image, count, data = _process_page(page, page_number, engine)
                ok, png = cv2.imencode(".png", cv2.cvtColor(image, cv2.COLOR_RGB2BGR))
                if not ok:
                    raise ValueError(f"Page {page_number}: could not create redacted page.")
                checkpoints.save_page(work_dir, page_number, png.tobytes(), count, data)
                del image, png
                gc.collect()
            else:
                _, count, data = saved
            redaction_count += count
            if data:
                extracted.append(data)
            else:
                unreadable.append(page_number)
            if on_page:
                on_page(page_number, len(document))
        with fitz.open() as redacted:
            for page_number, page in enumerate(document, 1):
                output_page = redacted.new_page(width=page.rect.width, height=page.rect.height)
                output_page.insert_image(output_page.rect, stream=checkpoints.page_image(work_dir, page_number))
            redacted.save(str(target), garbage=4, deflate=True)
    rows, warnings = merge_pages(extracted)
    if unreadable:
        warnings.append(f"Pages without a readable bill reference: {', '.join(map(str, unreadable))}.")
    if not rows:
        raise ValueError("No bill rows could be extracted. Please use a clearer scan.")
    return rows, redaction_count, warnings


def make_workbook(rows: list[dict]) -> bytes:
    if not isinstance(rows, list) or not rows:
        raise ValueError("Provide at least one reviewed bill row.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4B7A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34
    widths = (22, 16, 22, 16, 27, 25, 22, 25)
    for column, width in zip("ABCDEFGH", widths):
        sheet.column_dimensions[column].width = width
    for index, row in enumerate(rows, 2):
        if not isinstance(row, dict):
            raise ValueError(f"Row {index}: provide a bill record.")
        ref = str(row["bill_ref"]).strip().upper() if row.get("bill_ref") is not None else None
        if ref is not None and not REF_PATTERN.fullmatch(ref):
            raise ValueError(f"Row {index}: invalid bill reference.")
        hrn = str(row["hrn"]).strip() if row.get("hrn") is not None else None
        if hrn is not None and not re.fullmatch(r"[A-Za-z0-9-][A-Za-z0-9 /-]{0,39}", hrn):
            raise ValueError(f"Row {index}: invalid HRN.")
        try:
            bill_date = datetime.strptime(row["bill_date"], "%Y-%m-%d") if row.get("bill_date") else None
            visit_date = datetime.strptime(row["visit_date"], "%Y-%m-%d") if row.get("visit_date") else None
            total = float(row["total"]) if row.get("total") is not None else None
            medishield = float(row["medishield"]) if row.get("medishield") is not None else None
            medisave = float(row["medisave"]) if row.get("medisave") is not None else None
            cash = float(row["cash"]) if row.get("cash") is not None else None
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError(f"Row {index}: check dates and amounts.") from exc
        amounts = [value for value in (total, medishield, medisave, cash) if value is not None]
        if any(value < 0 for value in amounts):
            raise ValueError(f"Row {index}: amounts must be zero or greater.")
        if not all(math.isfinite(value) for value in amounts):
            raise ValueError(f"Row {index}: amounts must be finite numbers.")
        sheet.append((ref, bill_date, hrn, visit_date,
                      total, medishield, medisave, cash))
        for col in ("B", "D"):
            sheet[f"{col}{index}"].number_format = "dd mmm yyyy"
        for col in ("E", "F", "G", "H"):
            sheet[f"{col}{index}"].number_format = '#,##0.00'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{len(rows) + 1}"
    add_review_sheet(workbook, rows)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
