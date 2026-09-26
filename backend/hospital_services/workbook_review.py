"""Keep provisional OCR values visibly identified when a workbook is exported."""

from openpyxl.styles import Alignment, Font
from openpyxl.comments import Comment

from .merging import FIELD_NAMES


COLUMNS = {"bill_ref": "A", "bill_date": "B", "hrn": "C", "visit_date": "D",
           "total": "E", "medishield": "F", "medisave": "G", "cash": "H"}


def add_review_sheet(workbook, rows: list[dict]) -> None:
    flagged = [row for row in rows if row.get("review_fields") or row.get("review_notes")]
    if not flagged:
        return
    sheet = workbook.create_sheet("Review notes")
    sheet.append(("Bill Ref No.", "Field", "Displayed value", "OCR alternatives",
                  "Source", "Review reason"))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, 2):
        if not row.get("review_fields") and not row.get("review_notes"):
            continue
        fields = row.get("review_fields", [])
        candidates_by_field = row.get("field_candidates", {})
        pages = row.get("pages", [])
        if (not isinstance(fields, list) or len(fields) > len(FIELD_NAMES)
                or not isinstance(candidates_by_field, dict) or not isinstance(pages, list)):
            raise ValueError("Invalid OCR review metadata.")
        source = f"{row.get('source_file', '')} pages {', '.join(str(page) for page in pages)}".strip()
        for field in fields:
            if not isinstance(field, str) or field not in FIELD_NAMES:
                raise ValueError("Invalid review field.")
            candidates = candidates_by_field.get(field, [])
            if not isinstance(candidates, list) or len(candidates) > 100:
                raise ValueError("Invalid OCR alternatives.")
            alternatives = []
            for candidate in candidates:
                if not isinstance(candidate, dict) or not isinstance(candidate.get("pages", []), list):
                    raise ValueError("Invalid OCR alternative.")
                alternatives.append(f"{candidate.get('value', '')} (pages {', '.join(map(str, candidate.get('pages', [])))})")
            reason = ("Missing value" if row.get(field) is None else
                      "Conflicting readings" if len(candidates) > 1 else
                      "Low confidence or payment reconciliation; check source")
            values = (row.get("bill_ref"), FIELD_NAMES[field], row.get(field),
                      " / ".join(alternatives), source, reason)
            sheet.append(tuple("" if value is None else str(value)[:32767] for value in values))
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if field in COLUMNS:
                note = f"Provisional OCR value. {reason}.\n{source}\nAlternatives: {' / '.join(alternatives)}"
                workbook["Sheet1"][f"{COLUMNS[field]}{row_index}"].comment = Comment(note[:32767], "OCR review")
        notes = row.get("review_notes", [])
        if not isinstance(notes, list) or len(notes) > 100 or any(not isinstance(note, str) for note in notes):
            raise ValueError("Invalid invoice review notes.")
        for note in notes:
            if note.startswith("invoice pagination"):
                sheet.append((str(row.get("bill_ref") or ""), "Invoice pages", "", "", source[:32767], note[:32767]))
                for cell in sheet[sheet.max_row]:
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column, width in zip("ABCDEF", (23, 22, 24, 65, 50, 55)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
