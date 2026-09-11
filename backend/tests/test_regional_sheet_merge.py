import os
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from app import ExcelTransformer


HEADERS = [
    "No.",
    "Region",
    "Area",
    "PHPC Clinics?",
    "CLINIC",
    "Doctor",
    "Address",
    "Opening Hours 1",
    "Opening Hours 2",
    "Opening Hours 3",
    "Opening Hours 4",
    "Tel No.",
    "Fax No.",
    "Remarks",
    "Surcharge",
]


def add_regional_sheet(workbook, title, header_row, rows):
    worksheet = workbook.create_sheet(title)
    worksheet.cell(header_row - 1, 1, f"{title} Region")
    for column, value in enumerate(HEADERS, start=1):
        worksheet.cell(header_row, column, value)
    for row_number, values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row_number, column, value)


class RegionalSheetMergeTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.input_path = os.path.join(self.temp_dir.name, "regional.xlsx")

        workbook = Workbook()
        workbook.remove(workbook.active)
        add_regional_sheet(
            workbook,
            "NORTH",
            6,
            [
                [
                    1,
                    "NORTH REGION",
                    "WOODLANDS",
                    "YES",
                    "North Clinic",
                    "Dr North",
                    "1 NORTH ROAD SINGAPORE 123456",
                    "MON-FRI: 9AM-5PM",
                    "SAT: CLOSED",
                    "SUN: CLOSED",
                    "PH: CLOSED",
                    "61234567",
                    None,
                    None,
                    "NO",
                ]
            ],
        )
        add_regional_sheet(
            workbook,
            "EAST",
            3,
            [
                [
                    1,
                    "EAST REGION",
                    "TAMPINES",
                    "YES",
                    "East Clinic",
                    "Dr East",
                    "2 EAST ROAD SINGAPORE 654321",
                    "MON-FRI: 9AM-5PM",
                    "SAT: CLOSED",
                    "SUN: CLOSED",
                    "PH: CLOSED",
                    "68765432",
                    None,
                    None,
                    "NO",
                ]
            ],
        )
        add_regional_sheet(
            workbook,
            "MALAYSIA",
            4,
            [
                [
                    1,
                    "JOHOR",
                    "JOHOR BAHRU",
                    "NO",
                    "Malaysia Clinic",
                    "Dr South",
                    "3 JALAN TEST 80000 JOHOR MALAYSIA",
                    "MON-FRI: 9AM-5PM",
                    "SAT: CLOSED",
                    "SUN: CLOSED",
                    "PH: CLOSED",
                    "6071234567",
                    None,
                    None,
                    "NO",
                ]
            ],
        )
        workbook.save(self.input_path)
        workbook.close()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_detects_geographic_sheets_with_variable_header_rows(self):
        detection = ExcelTransformer.detect_regional_sheets(
            self.input_path,
            ["NORTH", "EAST", "MALAYSIA"],
        )

        self.assertEqual(
            [sheet["sheet_name"] for sheet in detection["valid_sheets"]],
            ["NORTH", "EAST", "MALAYSIA"],
        )
        self.assertEqual(
            [sheet["header_row"] for sheet in detection["valid_sheets"]],
            [5, 2, 3],
        )
        self.assertEqual(detection["invalid_sheets"], [])

    def test_combines_regional_rows_in_workbook_order(self):
        detection = ExcelTransformer.detect_regional_sheets(
            self.input_path,
            ["NORTH", "EAST", "MALAYSIA"],
        )

        combined, source_counts = ExcelTransformer.combine_regional_sheets(
            self.input_path,
            detection["valid_sheets"],
        )

        self.assertEqual(list(combined["CLINIC"]), [
            "North Clinic",
            "East Clinic",
            "Malaysia Clinic",
        ])
        self.assertEqual(
            source_counts,
            {"NORTH": 1, "EAST": 1, "MALAYSIA": 1},
        )

    def test_canonicalizes_compatible_header_case_before_concat(self):
        workbook = load_workbook(self.input_path)
        east_sheet = workbook["EAST"]
        east_sheet.cell(3, 5, "Clinic")
        east_sheet.cell(3, 7, "ADDRESS")
        workbook.save(self.input_path)
        workbook.close()

        detection = ExcelTransformer.detect_regional_sheets(
            self.input_path,
            ["NORTH", "EAST", "MALAYSIA"],
        )
        combined, _ = ExcelTransformer.combine_regional_sheets(
            self.input_path,
            detection["valid_sheets"],
        )

        self.assertEqual(detection["invalid_sheets"], [])
        self.assertIn("CLINIC", combined.columns)
        self.assertNotIn("Clinic", combined.columns)
        self.assertIn("Address", combined.columns)
        self.assertNotIn("ADDRESS", combined.columns)
        self.assertEqual(
            list(combined["CLINIC"]),
            ["North Clinic", "East Clinic", "Malaysia Clinic"],
        )

    def test_counts_duplicates_without_removing_records(self):
        dataframe = pd.DataFrame(
            {
                "Name": ["Same Clinic", " same   clinic ", "Other Clinic"],
                "Address1": ["1 Test Road", " 1 TEST ROAD ", "2 Test Road"],
            }
        )

        duplicate_count = ExcelTransformer.count_possible_duplicate_clinics(
            dataframe
        )

        self.assertEqual(duplicate_count, 1)
        self.assertEqual(len(dataframe), 3)

    def test_maps_income_regional_opening_hour_columns(self):
        column_map = ExcelTransformer.map_columns(HEADERS)

        self.assertEqual(column_map["mon_fri_am"], "Opening Hours 1")
        self.assertEqual(column_map["sat_simple"], "Opening Hours 2")
        self.assertEqual(column_map["sun_simple"], "Opening Hours 3")
        self.assertEqual(column_map["holiday_simple"], "Opening Hours 4")

    def test_rejects_partial_merge_when_a_regional_sheet_is_invalid(self):
        workbook = load_workbook(self.input_path)
        worksheet = workbook.create_sheet("WEST")
        worksheet.append(["Summary only", "No clinic table here"])
        workbook.save(self.input_path)
        workbook.close()

        output_dir = os.path.join(self.temp_dir.name, "invalid-output")
        os.makedirs(output_dir)
        with patch.object(ExcelTransformer, "transform_sheet") as transform_sheet:
            result = ExcelTransformer.transform_excel_multi_sheet(
                self.input_path,
                output_dir,
                "job-invalid",
                use_google_api=False,
            )

        self.assertFalse(result["success"])
        self.assertIn("Regional sheet validation failed", result["message"])
        self.assertIn("WEST", result["message"])
        transform_sheet.assert_not_called()

    def test_multi_sheet_transform_splits_singapore_and_malaysia_outputs(self):
        transformed = pd.DataFrame(
            {
                "Code": ["AUTO_0001", "AUTO_0002", "AUTO_0003"],
                "Name": ["North Clinic", "East Clinic", "Malaysia Clinic"],
                "Address1": [
                    "1 NORTH ROAD SINGAPORE 123456",
                    "2 EAST ROAD SINGAPORE 654321",
                    "3 JALAN TEST 80000 JOHOR MALAYSIA",
                ],
                "PostalCode": ["123456", "654321", "80000"],
                "Country": ["SINGAPORE", "SINGAPORE", "MALAYSIA"],
                "Latitude": [1.30, 1.31, None],
                "Longitude": [103.80, 103.81, None],
            }
        )

        def fake_transform(path, sheet_name, terminated_ids, use_google_api):
            self.assertEqual(sheet_name, "List")
            combined_source = pd.read_excel(path, sheet_name="List")
            self.assertEqual(len(combined_source), 3)
            return {
                "success": True,
                "dataframe": transformed,
                "geocoding_methods": ["postal_code", "address", None],
                "records_processed": 3,
                "terminated_clinics_filtered": 0,
                "filtered_provider_codes": [],
                "geocoding_stats": {
                    "total_records": 3,
                    "successful_geocodes": 2,
                    "postal_code_matches": 1,
                    "address_geocodes": 1,
                    "failed_geocodes": 1,
                    "success_rate": "66.7%",
                },
            }

        output_dir = os.path.join(self.temp_dir.name, "output")
        os.makedirs(output_dir)
        with patch.object(
            ExcelTransformer,
            "transform_sheet",
            side_effect=fake_transform,
        ):
            result = ExcelTransformer.transform_excel_multi_sheet(
                self.input_path,
                output_dir,
                "job-123",
                use_google_api=False,
            )

        self.assertTrue(result["success"])
        self.assertEqual(
            result["output_files"],
            [
                "job-123_Singapore.xlsx",
                "job-123_Malaysia.xlsx",
            ],
        )
        self.assertEqual(
            [item["sheet_name"] for item in result["results"]],
            ["SINGAPORE", "MALAYSIA"],
        )
        self.assertEqual(
            [item["records_processed"] for item in result["results"]],
            [2, 1],
        )
        self.assertEqual(
            result["results"][0]["geocoding_stats"],
            {
                "total_records": 2,
                "successful_geocodes": 2,
                "postal_code_matches": 1,
                "address_geocodes": 1,
                "failed_geocodes": 0,
                "success_rate": "100.0%",
            },
        )
        self.assertEqual(
            result["results"][1]["geocoding_stats"],
            {
                "total_records": 1,
                "successful_geocodes": 0,
                "postal_code_matches": 0,
                "address_geocodes": 0,
                "failed_geocodes": 1,
                "success_rate": "0.0%",
            },
        )
        self.assertEqual(result["regional_merge"]["source_sheet_count"], 3)
        self.assertTrue(result["regional_merge"]["split_by_country"])
        self.assertEqual(
            result["regional_merge"]["output_countries"],
            ["SINGAPORE", "MALAYSIA"],
        )
        self.assertEqual(
            result["regional_merge"]["source_sheets"],
            ["NORTH", "EAST", "MALAYSIA"],
        )
        expected_rows = {
            "job-123_Singapore.xlsx": 2,
            "job-123_Malaysia.xlsx": 1,
        }
        for filename, row_count in expected_rows.items():
            output_path = os.path.join(output_dir, filename)
            output_workbook = load_workbook(output_path, read_only=True)
            try:
                self.assertEqual(output_workbook.sheetnames, ["List"])
                self.assertEqual(output_workbook["List"].max_row - 1, row_count)
            finally:
                output_workbook.close()


if __name__ == "__main__":
    unittest.main()
