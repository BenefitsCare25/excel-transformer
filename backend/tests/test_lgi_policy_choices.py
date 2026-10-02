"""Employee choices must change every report and its totals consistently."""
from datetime import datetime
import unittest

import openpyxl

from flex_services.companies.lgi import run, validate_inputs
from flex_services.errors import FlexInputError
from tests import test_lgi


class LGIPolicyChoiceTests(unittest.TestCase):
    def setUp(self):
        self.fixture = test_lgi.LGIClaimsTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        # Two claimants, both with records before the policy start.
        second = self.fixture.listing.iloc[0].copy()
        second['Employee ID No.'] = 'LCM-002'
        second['Employee Name'] = 'Second Employee'
        self.fixture.listing.loc[1] = second
        self.fixture.claims.loc[1, ['Staff ID', 'Employee Name']] = ['LCM-002', 'Second Employee']
        second_util = self.fixture.utilization.iloc[0].copy()
        second_util['Staff ID'] = 'LCM-002'
        second_util['Employee Name'] = 'Second Employee'
        self.fixture.utilization.loc[1] = second_util
        self.fixture.prepare_utilization()
        self.fixture.utilization['Benefit Start Date'] = datetime(2024, 10, 1)
        self.fixture.utilization.to_excel(self.fixture.root / 'utilization.xlsx', index=False)
        self.files = {key: self.fixture.root / f'{key}.xlsx' for key in ['claims', 'listing', 'utilization']}

    def test_each_employee_requires_a_choice(self):
        for decisions in [None, {'LCM-001': 'include'}]:
            with self.assertRaises(FlexInputError) as caught:
                validate_inputs(self.files, '2026-07-01', decisions)
            self.assertEqual(len(caught.exception.feedback['policy_review']), 2)

    def test_include_exclude_and_mixed_choices_update_all_workbooks(self):
        for decisions, ids, amount in [
            ({'LCM-001': 'include', 'LCM-002': 'include'}, ['LCM-001', 'LCM-002'], 160.3),
            ({'LCM-001': 'include', 'LCM-002': 'exclude'}, ['LCM-001'], 120.1),
            ({'LCM-001': 'exclude', 'LCM-002': 'include'}, ['LCM-002'], 40.2),
            ({'LCM-001': 'exclude', 'LCM-002': 'exclude'}, [], 0),
        ]:
            with self.subTest(decisions=decisions):
                result = run(self.files, '2026-07-01', self.fixture.root / 'out', decisions)
                self.assertEqual(result['grand_total'], amount)
                self.assertEqual(result['employees'], len(ids))
                for index, id_column, first_row, money_columns in [
                    (0, 3, 2, [12]), (1, 1, 2, [6, 7]), (2, 2, 7, [10]),
                ]:
                    workbook = openpyxl.load_workbook(result['outputs'][index], data_only=True)
                    try:
                        sheet = workbook.active
                        self.assertEqual([sheet.cell(row, id_column).value for row in range(first_row, sheet.max_row + 1)], ids)
                        total = sum(sheet.cell(row, col).value or 0 for row in range(first_row, sheet.max_row + 1) for col in money_columns)
                        self.assertAlmostEqual(total, amount)
                    finally:
                        workbook.close()

    def test_inclusion_still_checks_balances_and_employee_details(self):
        self.fixture.utilization.loc[0, 'Balance Available Allocation Amt'] = 999
        self.fixture.utilization.to_excel(self.files['utilization'], index=False)
        with self.assertRaisesRegex(FlexInputError, 'does not reconcile'):
            validate_inputs(self.files, '2026-07-01', {'LCM-001': 'include', 'LCM-002': 'exclude'})

    def test_missing_future_and_ambiguous_records_cannot_be_included(self):
        for kind in ['missing', 'future', 'duplicate']:
            frame = self.fixture.utilization.copy()
            if kind == 'missing':
                frame = frame.loc[frame['Staff ID'].ne('LCM-001')]
            elif kind == 'future':
                frame.loc[0, 'Benefit Start Date'] = datetime(2026, 8, 1)
            else:
                frame.loc[2] = frame.iloc[0].copy()
            frame.to_excel(self.files['utilization'], index=False)
            with self.assertRaises(FlexInputError):
                validate_inputs(self.files, '2026-07-01', {'LCM-001': 'include', 'LCM-002': 'include'})


if __name__ == '__main__':
    unittest.main()
