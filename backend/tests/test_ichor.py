"""Ichor claim-category and summary layout checks using synthetic data."""

from datetime import datetime
from pathlib import Path
import tempfile
import unittest

import openpyxl
import pandas as pd

from flex_services.companies.ichor import run
from flex_services.companies.ichor.constants import (
    ICHOR_ENTITY,
    SUMMARY_COLUMNS,
    SUMMARY_GROUPS,
    WELLNESS,
)


class IchorClaimsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.claims = pd.DataFrame(
            [
                {
                    "Entity": ICHOR_ENTITY,
                    "Staff ID": 123456,
                    "Employee Name": "Sample Employee",
                    "Reference No.": "ICHOR-000001",
                    "Claim Type": "Polyclinic :: Dependent",
                    "Incurred Date": datetime(2026, 9, 1),
                    "Service Provider": "Sample Clinic",
                    "Converted Currency": "SGD",
                    "Converted Incurred Amt": 40.67,
                    "Payment Amt": 40.67,
                    "Status": "Approved",
                    "Paid Date": datetime(2026, 9, 28),
                    "Admin Remark": "",
                },
                {
                    "Entity": ICHOR_ENTITY,
                    "Staff ID": 123456,
                    "Employee Name": "Sample Employee",
                    "Reference No.": "ICHOR-000002",
                    "Claim Type": f"{WELLNESS} :: Employee",
                    "Incurred Date": datetime(2026, 9, 2),
                    "Service Provider": "Sample Gym",
                    "Converted Currency": "SGD",
                    "Converted Incurred Amt": 100.00,
                    "Payment Amt": 100.00,
                    "Status": "Approved",
                    "Paid Date": datetime(2026, 9, 28),
                    "Admin Remark": "",
                },
            ]
        )
        self.listing = pd.DataFrame(
            [
                {
                    "User ID": 123456,
                    "Employee Name": "Sample Employee",
                    "Last Day of Service": None,
                }
            ]
        )

    def test_new_categories_are_supported_and_grouped_in_summary(self):
        claims_path = self.root / "claims.xlsx"
        listing_path = self.root / "listing.xlsx"
        self.claims.to_excel(claims_path, index=False)
        self.listing.to_excel(listing_path, index=False)

        result = run(
            {"claims": claims_path, "listing": listing_path},
            "2026-09-01",
            self.root / "output",
        )

        self.assertEqual(result["errors"], 0)
        self.assertEqual(result["grand_total"], 140.67)
        summary = openpyxl.load_workbook(result["outputs"][1], data_only=True)
        self.addCleanup(summary.close)
        sheet = summary.active

        polyclinic_column = 3 + SUMMARY_COLUMNS.index(("Polyclinic", "Dependent"))
        wellness_column = 3 + SUMMARY_COLUMNS.index((WELLNESS, "Employee"))
        grand_total_column = 3 + len(SUMMARY_COLUMNS)
        self.assertEqual(sheet.cell(5, polyclinic_column).value, 40.67)
        self.assertEqual(sheet.cell(5, wellness_column).value, 100)
        self.assertEqual(sheet.cell(5, grand_total_column).value, 140.67)
        self.assertEqual(sheet.cell(5, grand_total_column + 1).value, 40.67)
        self.assertEqual(sheet.cell(5, grand_total_column + 2).value, 0)
        self.assertEqual(sheet.cell(5, grand_total_column + 3).value, 100)
        self.assertEqual(
            sheet.max_column,
            grand_total_column + len(SUMMARY_GROUPS),
        )


if __name__ == "__main__":
    unittest.main()
