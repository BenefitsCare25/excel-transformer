"""LGI report checks using synthetic employee data only."""

from datetime import datetime
from pathlib import Path
import tempfile
import unittest

import openpyxl
import pandas as pd

from flex_services.errors import FlexInputError
from flex_services.companies.lgi.constants import ENTITY, DETAIL_HEADERS, SUMMARY_HEADERS
from flex_services.companies.lgi.processing import (
    balance_validations, claim_rule, employment_validations,
    load_claims, load_listing, leaver_validations, load_utilization, policy_period,
)
from flex_services.companies.lgi.workbooks import write_details, write_summary, write_utilization
from flex_services.companies.lgi import run


class LGIClaimsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.listing = pd.DataFrame([{
            'Entity': ENTITY, 'Employee ID No.': 'LCM-001', 'Employee Name': 'Sample Employee',
            'Designation': 'Manager', 'Cost Centre': '012', 'Department': 'Operations',
            'Date of Hire': datetime(2024, 1, 1), 'Last Day of Service': datetime(2026, 7, 31),
        }])
        self.claims = pd.DataFrame([{
            'Entity': ENTITY, 'Staff ID': 'LCM-001', 'Employee Name': 'Sample Employee',
            'Claimant Name': 'Sample Child', 'Relation': 'Child', 'Reference No.': 'LGI-000001',
            'Claim Type': 'Dental', 'TAX': 'No', 'CPF': 'No',
            'Incurred Date': datetime(2026, 6, 15), 'Service Provider': '=1+1',
            'Converted Currency': 'SGD', 'Converted Incurred Amt': 180.55,
            'Payment Amt': 120.10, 'Status': 'Approved', 'Paid Date': datetime(2026, 7, 25),
            'Admin Remark': 'Partially reimbursed',
        }, {
            'Entity': ENTITY, 'Staff ID': 'LCM-001', 'Employee Name': 'Sample Employee',
            'Claimant Name': 'Sample Employee', 'Relation': 'Self', 'Reference No.': 'LGI-000002',
            'Claim Type': 'Holiday travel insurance, admission fees to local attraction etc',
            'TAX': 'Yes', 'CPF': 'Yes', 'Incurred Date': datetime(2026, 6, 20),
            'Service Provider': 'Sample Insurer', 'Converted Currency': 'SGD',
            'Converted Incurred Amt': 40.20, 'Payment Amt': 40.20,
            'Status': 'Approved', 'Paid Date': datetime(2026, 7, 25), 'Admin Remark': '',
        }])
        self.utilization = pd.DataFrame([{
            'Entity': ENTITY, 'Staff ID': 'LCM-001', 'Employee Name': 'Sample Employee',
            'Date of Hire': datetime(2024, 1, 1), 'Last Day of Service': None,
            'Benefit Start Date': datetime(2025, 10, 1), 'Benefit End Date': None,
            'Wallet': 'FSA', 'Total Allocation Amt': 200., 'Buy Leave Amt': 20.,
            'Sell Leave Amt': 100., 'Selection Amt': 50., 'Deals Amt': 5.,
            'Claims Payment Amt': 300., 'Total Utilized Amt (L-M+N+O+P)': 275.,
            'Pending Claims Payment Amt': 77., 'Balance Available Allocation Amt': -75.,
        }])

    def prepare(self, month='2026-07-01'):
        listing_path = self.root / 'listing.xlsx'
        claims_path = self.root / 'claims.xlsx'
        self.listing.to_excel(listing_path, index=False)
        # openpyxl keeps the synthetic provider as literal text in the input, too.
        self.claims.to_excel(claims_path, index=False, engine='openpyxl')
        workbook = openpyxl.load_workbook(claims_path)
        provider_col = list(self.claims.columns).index('Service Provider') + 1
        workbook.active.cell(2, provider_col).data_type = 's'
        workbook.save(claims_path)
        workbook.close()
        return load_claims(claims_path, load_listing(listing_path), month)

    def test_reference_layout_and_per_claim_summary(self):
        claims = self.prepare()
        month = datetime(2026, 7, 1)
        details = openpyxl.load_workbook(write_details(claims, month, self.root))
        summary = openpyxl.load_workbook(write_summary(claims, month, self.root))
        self.addCleanup(details.close)
        self.addCleanup(summary.close)
        detail = details.active
        sheet = summary.active
        self.assertEqual(tuple(c.value for c in detail[1]), DETAIL_HEADERS)
        self.assertEqual(tuple(c.value for c in sheet[1]), SUMMARY_HEADERS)
        self.assertEqual((sheet.max_row, detail.max_row), (3, 3))
        self.assertEqual([sheet.cell(r, 6).value for r in [2, 3]], [0, 40.2])
        self.assertEqual([sheet.cell(r, 7).value for r in [2, 3]], [120.1, 0])
        self.assertEqual([sheet.cell(r, 8).value for r in [2, 3]], ['No', 'Yes'])
        self.assertEqual(sheet['C2'].value, 'Manager')
        self.assertEqual(sheet['D2'].value, '012')
        self.assertEqual(detail['K2'].value, 180.55)
        self.assertEqual(detail['L2'].value, 120.1)
        self.assertEqual(detail['N2'].value, 'Paid')
        self.assertEqual(detail['O2'].value, 'Partially reimbursed')
        self.assertEqual(detail['M2'].value, '=1+1')
        self.assertEqual(detail['M2'].data_type, 's')
        self.assertEqual(detail['B2'].number_format, 'dd/mm/yyyy')
        self.assertEqual(len(leaver_validations(claims)), 1)
        self.assertAlmostEqual(leaver_validations(claims)[0]['amount'], 160.3)

    def test_latest_checklist_categories_and_short_labels(self):
        expected = {
            "Holiday travel insurance, admission fees to local attraction etc": ("Yes", "Yes"),
            "Children's Education Tuition Fees": ("Yes", "Yes"),
            "Spa/Wellness Services": ("Yes", "Yes"),
            "Medical Appliances": ("Yes", "Yes"),
            "Fertility Treatment": ("No", "Yes"),
            "Lasik Surgery": ("No", "Yes"),
            "Self-Improvement Course Fees": ("Yes", "No"),
        }
        for label, rule in expected.items():
            with self.subTest(label=label):
                self.assertEqual(claim_rule(label), rule)
                self.assertEqual(claim_rule(f"  {label.upper()}  "), rule)

        self.claims.loc[0, ['Claim Type', 'TAX', 'CPF']] = [
            'Fertility Treatment', 'No', 'Yes']
        prepared = self.prepare()
        self.assertEqual(prepared.loc[0, 'Claim Type'], 'Fertility Treatment')
        self.assertEqual(prepared.loc[0, 'ClaimRule'], ('No', 'Yes'))

    def test_tax_and_cpf_combinations_remain_independent(self):
        self.claims.loc[0, ['Claim Type', 'TAX', 'CPF']] = [
            'Family Holidays (hotel, chalets, holiday bungalows, tour package, air tickets)', 'Yes', 'No']
        self.claims.loc[1, ['Claim Type', 'TAX', 'CPF']] = ['Maternity', 'No', 'Yes']
        result = self.prepare()
        self.assertEqual(result['Taxable'].tolist(), [120.1, 0])
        self.assertEqual(result['Non Taxable'].tolist(), [0, 40.2])
        self.assertEqual(result['CPF'].tolist(), ['No', 'Yes'])

    def test_unknown_claim_and_conflicting_flags_block(self):
        for col, value, expected in [('Claim Type', 'Unknown category', 'unsupported Claim Type'),
                                     ('TAX', 'Yes', 'TAX conflicts'), ('CPF', 'Yes', 'CPF conflicts')]:
            with self.subTest(column=col):
                original = self.claims.loc[0, col]
                self.claims.loc[0, col] = value
                with self.assertRaisesRegex(FlexInputError, expected):
                    self.prepare()
                self.claims.loc[0, col] = original

    def test_invalid_claim_rows_block(self):
        for col, value, expected in [
            ('Payment Amt', -1, 'negative Payment Amt'),
            ('Payment Amt', float('inf'), 'invalid or missing Payment Amt'),
            ('Payment Amt', 181, 'exceeds Converted'),
            ('Status', 'Pending', 'Approved claims only'),
            ('Converted Currency', 'USD', 'must be in SGD'),
            ('Reference No.', 'LGI-000002', 'duplicate Reference'),
            ('Staff ID', 'LCM-MISSING', 'missing from listing'),
            ('Employee Name', 'Different Name', 'name or entity mismatch'),
        ]:
            with self.subTest(column=col, value=value):
                original = self.claims.loc[0, col]
                self.claims.loc[0, col] = value
                with self.assertRaisesRegex(FlexInputError, expected):
                    self.prepare()
                self.claims.loc[0, col] = original

    def test_wrong_month_blocks(self):
        with self.assertRaisesRegex(FlexInputError, 'outside 2026-08'):
            self.prepare('2026-08-01')

    def test_missing_columns_and_duplicate_listing_block(self):
        self.claims = self.claims.drop(columns='CPF')
        with self.assertRaisesRegex(FlexInputError, 'missing required column.*CPF'):
            self.prepare()
        self.listing = pd.concat([self.listing, self.listing], ignore_index=True)
        with self.assertRaisesRegex(FlexInputError, 'duplicate Employee ID'):
            self.prepare()

    def prepare_utilization(self):
        claims = self.prepare()
        path = self.root / 'utilization.xlsx'
        self.utilization.to_excel(path, index=False)
        listing = load_listing(self.root / 'listing.xlsx')
        return load_utilization(path, listing, claims, '2026-07-01')

    def test_utilization_arithmetic_and_source_balances(self):
        frame, excluded = self.prepare_utilization()
        self.assertEqual(excluded, 0)
        row = frame.iloc[0]
        self.assertEqual(row['FlexUsed'], -25.)
        self.assertEqual(row['TransferredFSA'], 225.)
        self.assertEqual(row['MonthlyClaims'], 160.3)
        self.assertEqual(len(balance_validations(frame)), 1)
        self.assertEqual(len(employment_validations(frame)), 1)
        path = write_utilization(frame, datetime(2026, 7, 1), self.root)
        workbook = openpyxl.load_workbook(path, data_only=True)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet.title, '1 Oct 2025  to 30 Sep 2026')
        self.assertEqual(sheet['G7'].value, 200.)
        self.assertEqual(sheet['H7'].value, -25.)
        self.assertEqual(sheet['I7'].value, 225.)
        self.assertEqual(sheet['J7'].value, 160.3)
        self.assertEqual(sheet['K7'].value, 300.)
        self.assertIsNone(sheet['L7'].value)
        self.assertEqual(sheet['M7'].value, -75.)
        self.assertEqual(sheet['N7'].value, datetime(2026, 7, 31))
        self.assertEqual(sheet['D7'].value, 'Manager')

    def test_policy_filter_preserves_zero_allocation_leavers(self):
        for staff, benefit, last_day in [
            ('LCM-002', datetime(2025, 10, 1), datetime(2026, 1, 1)),
            ('LCM-003', datetime(2024, 10, 1), datetime(2025, 1, 1)),
            ('LCM-004', datetime(2026, 8, 1), None),
        ]:
            listing_row = self.listing.iloc[0].to_dict()
            listing_row.update({'Employee ID No.': staff, 'Last Day of Service': last_day})
            self.listing.loc[len(self.listing)] = listing_row
            util_row = self.utilization.iloc[0].to_dict()
            util_row.update({'Staff ID': staff, 'Benefit Start Date': benefit,
                             'Last Day of Service': last_day})
            for col in self.utilization.columns:
                if 'Amt' in col:
                    util_row[col] = 0.
            self.utilization.loc[len(self.utilization)] = util_row
        frame, excluded = self.prepare_utilization()
        self.assertEqual(frame['Staff ID'].tolist(), ['LCM-001', 'LCM-002'])
        self.assertEqual(excluded, 2)
        self.assertEqual(frame.iloc[1]['Total Allocation Amt'], 0.)
        self.assertEqual(frame.iloc[1]['MonthlyClaims'], 0.)

    def test_prior_and_current_policy_rows_for_same_employee_are_allowed(self):
        historical = self.utilization.iloc[0].copy()
        historical['Benefit Start Date'] = datetime(2024, 10, 1)
        historical['Last Day of Service'] = datetime(2025, 9, 30)
        self.utilization = pd.concat(
            [pd.DataFrame([historical]), self.utilization], ignore_index=True)

        frame, excluded = self.prepare_utilization()

        self.assertEqual(excluded, 1)
        self.assertEqual(frame['Staff ID'].tolist(), ['LCM-001'])
        self.assertEqual(frame.iloc[0]['Benefit Start Date'], datetime(2025, 10, 1))

    def test_utilization_inconsistent_balances_block(self):
        for col in ['Balance Available Allocation Amt', 'Total Utilized Amt (L-M+N+O+P)']:
            with self.subTest(column=col):
                original = self.utilization.loc[0, col]
                self.utilization.loc[0, col] = 999.
                with self.assertRaisesRegex(FlexInputError, 'does not reconcile'):
                    self.prepare_utilization()
                self.utilization.loc[0, col] = original

    def test_utilization_older_than_claims_blocks(self):
        self.utilization.loc[0, ['Claims Payment Amt', 'Total Utilized Amt (L-M+N+O+P)',
                                 'Balance Available Allocation Amt']] = [100., 75., 125.]
        with self.assertRaisesRegex(FlexInputError, 'less than this month'):
            self.prepare_utilization()

    def test_utilization_missing_claimant_and_duplicate_wallet_block(self):
        self.utilization = pd.concat([self.utilization, self.utilization], ignore_index=True)
        with self.assertRaisesRegex(FlexInputError, 'duplicate Staff ID'):
            self.prepare_utilization()
        self.utilization = self.utilization.iloc[:1].copy()
        self.utilization.loc[0, 'Staff ID'] = 'LCM-002'
        listing_row = self.listing.iloc[0].to_dict()
        listing_row['Employee ID No.'] = 'LCM-002'
        self.listing.loc[len(self.listing)] = listing_row
        with self.assertRaisesRegex(FlexInputError, 'missing from current policy'):
            self.prepare_utilization()

    def test_policy_year_rolls_in_october(self):
        self.assertEqual(policy_period('2026-09-01'),
                         (pd.Timestamp(2025, 10, 1), pd.Timestamp(2026, 9, 30)))
        self.assertEqual(policy_period('2026-10-01'),
                         (pd.Timestamp(2026, 10, 1), pd.Timestamp(2027, 9, 30)))

    def test_complete_adapter_writes_three_unencrypted_workbooks(self):
        self.prepare_utilization()
        files = {key: self.root / f'{key}.xlsx' for key in ['claims', 'listing', 'utilization']}
        result = run(files, '2026-07-01', self.root / 'out')
        self.assertEqual(len(result['outputs']), 3)
        self.assertEqual(result['errors'], 0)
        self.assertEqual(result['grand_total'], 160.3)
        self.assertEqual(result['breakdown_rows'], 2)
        self.assertEqual(result['employees'], 1)
        for path in result['outputs']:
            workbook = openpyxl.load_workbook(path)
            self.assertFalse(workbook.active.protection.sheet)
            workbook.close()

    def test_bad_utilization_creates_no_partial_reports(self):
        self.prepare_utilization()
        self.utilization.loc[0, 'Balance Available Allocation Amt'] = 999.
        self.utilization.to_excel(self.root / 'utilization.xlsx', index=False)
        files = {key: self.root / f'{key}.xlsx' for key in ['claims', 'listing', 'utilization']}
        with self.assertRaises(FlexInputError):
            run(files, '2026-07-01', self.root / 'out')
        self.assertFalse((self.root / 'out').exists())


if __name__ == '__main__':
    unittest.main()
