import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl
import pandas as pd

from flex_services.companies.stm import (
    DEFAULT_IT15_TEMPLATE,
    LEAVERS_COLUMNS,
    PAYROLL_WAGE_TYPE_LOOKUP_SHEET,
    _ensure_payroll_wage_type_lookup,
    claim_type_wage_code,
    run,
)


class StmClaimTypeMappingTests(unittest.TestCase):
    def test_maps_supported_claim_types_to_their_wage_codes(self):
        expected_codes = {
            'Optical': 'Optical(T)',
            'Childcare': 'HealthS/ChildC(NT)',
            'Health Screening': 'HealthS/ChildC(NT)',
            'Dental': 'N-Medi/Dental(NT)',
            'Medical-related': 'N-Medi/Dental(NT)',
            'Outpatient GP (capped at S$30 per visit)': 'N-Medi/Dental(NT)',
            'Outpatient Medical expenses at Singapore  Government Polyclinics': 'N-Medi/Dental(NT)',
            'Gym/Fitness Membership': 'Gym/Fitness(T&C)',
            'Alternative Treatment (Chiropractic)': 'Chiropractic(T&C)',
        }

        for claim_type, expected_code in expected_codes.items():
            with self.subTest(claim_type=claim_type):
                self.assertEqual(claim_type_wage_code(claim_type), expected_code)

    def test_normalises_claim_type_case_and_whitespace(self):
        self.assertEqual(
            claim_type_wage_code('  GYM/FITNESS   MEMBERSHIP  '),
            'Gym/Fitness(T&C)',
        )

    def test_rejects_unmapped_claim_type(self):
        self.assertIsNone(claim_type_wage_code('Unapproved wellness benefit'))

    def test_updates_legacy_it15_lookup_with_confirmed_wage_types(self):
        workbook = openpyxl.load_workbook(DEFAULT_IT15_TEMPLATE)
        try:
            lookup_end_row = _ensure_payroll_wage_type_lookup(workbook)
            # A second call proves retries do not duplicate lookup rows.
            self.assertEqual(_ensure_payroll_wage_type_lookup(workbook), lookup_end_row)

            worksheet = workbook[PAYROLL_WAGE_TYPE_LOOKUP_SHEET]
            mappings = {
                worksheet.cell(row, 9).value: worksheet.cell(row, 10).value
                for row in range(5, lookup_end_row + 1)
            }
            self.assertEqual(mappings['Chiropractic(T&C)'], 2234)
            self.assertEqual(mappings['Gym/Fitness(T&C)'], 2235)
            self.assertEqual(
                [
                    worksheet.cell(row, 9).value
                    for row in range(5, lookup_end_row + 1)
                ].count('Gym/Fitness(T&C)'),
                1,
            )
            self.assertEqual(worksheet.tables['Table3'].ref, f'H4:P{lookup_end_row}')
        finally:
            workbook.close()

    def test_generated_it15_uses_confirmed_wage_types_and_numeric_codes(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            entity = 'STMICROELECTRONICS ASIA PACIFIC PTE LTD - G0005086'
            claims = pd.DataFrame([
                {
                    'Staff ID': '123456',
                    'Employee Name': 'Payroll, Test',
                    'Reference No.': 'STM-GYM',
                    'Entity': entity,
                    'Claim Type': 'Gym/Fitness Membership',
                    'Payment Amt': 100.00,
                    'Status': 'Approved',
                    'Incurred Date': pd.Timestamp('2026-09-01'),
                },
                {
                    'Staff ID': '123456',
                    'Employee Name': 'Payroll, Test',
                    'Reference No.': 'STM-CHIRO',
                    'Entity': entity,
                    'Claim Type': 'Alternative Treatment (Chiropractic)',
                    'Payment Amt': 200.00,
                    'Status': 'Approved',
                    'Incurred Date': pd.Timestamp('2026-09-02'),
                },
            ])
            listing = pd.DataFrame([{
                'User ID': '123456',
                'Cost Centre': 'SH1234',
                'Location Description': 'Singapore',
                'Last Day of Service': pd.NaT,
            }])
            leaver = dict.fromkeys(LEAVERS_COLUMNS)
            leaver.update({
                'EmpID': '999999',
                'Name': 'Unrelated, Leaver',
                'EmailDate': pd.Timestamp('2026-08-01'),
            })

            claims_path = root / 'claims.xlsx'
            listing_path = root / 'listing.xlsx'
            leavers_path = root / 'leavers.xlsx'
            output_dir = root / 'output'
            claims.to_excel(claims_path, index=False)
            listing.to_excel(listing_path, index=False)
            pd.DataFrame([leaver]).to_excel(leavers_path, index=False, startrow=1)
            output_dir.mkdir()

            result = run(
                {
                    'claims': claims_path,
                    'listing': listing_path,
                    'leavers': leavers_path,
                },
                '2026-09-01',
                output_dir,
            )
            payroll_path = next(
                Path(path) for path in result['outputs']
                if 'Payroll_Additional' in Path(path).name
            )
            workbook = openpyxl.load_workbook(payroll_path, data_only=False)
            try:
                worksheet = workbook['SG1xPaymtTemplate']
                self.assertEqual(worksheet['H20'].value, 'Chiropractic(T&C)')
                self.assertEqual(worksheet['H21'].value, 'Gym/Fitness(T&C)')
                self.assertIn('$I$4:$J$10', worksheet['R20'].value)
                self.assertIn('$I$4:$J$10', worksheet['R21'].value)

                lookup = workbook[PAYROLL_WAGE_TYPE_LOOKUP_SHEET]
                lookup_codes = {
                    lookup.cell(row, 9).value: lookup.cell(row, 10).value
                    for row in range(5, 11)
                }
                self.assertEqual(lookup_codes['Chiropractic(T&C)'], 2234)
                self.assertEqual(lookup_codes['Gym/Fitness(T&C)'], 2235)
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()
