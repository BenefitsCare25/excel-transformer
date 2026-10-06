import os
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from app import ExcelTransformer


class TransformerCodeTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.input_path = os.path.join(self.temp_dir.name, "panels.xlsx")
        self.source = pd.DataFrame({
            "S/N": [1, 2, 3, 4],
            "Clinic ID": ["P100", None, "P100", "P100-1"],
            "Clinic Name": ["Alpha Clinic", "Beta Clinic", "Gamma Clinic", "Delta Clinic"],
            "Region": ["NORTH"] * 4,
            "Area": ["TEST AREA"] * 4,
            "Address": ["1 TEST ROAD SINGAPORE 012345"] * 4,
            "Postal Code": ["012345"] * 4,
            "Tel No.": ["61234567"] * 4,
        })
        geocoding = patch("app.GeocodingService")
        self.geocoder = geocoding.start().return_value
        self.addCleanup(geocoding.stop)
        self.geocoder.geocode.return_value = (None, None, None)
        self.geocoder.get_stats.return_value = {}

    def transform(self, source=None, terminated_ids=None):
        if source is None:
            source = self.source
        source.to_excel(self.input_path, sheet_name="Dental Panel", index=False)
        result = ExcelTransformer.transform_sheet(
            self.input_path, "Dental Panel", terminated_ids, use_google_api=False,
        )
        self.assertTrue(result["success"], result.get("error_details", result["message"]))
        return result

    def assert_codes(self, dataframe):
        self.assertEqual(dataframe.columns[0], "Code")
        self.assertEqual(dataframe["Code"].tolist(), list(range(1, len(dataframe) + 1)))
        self.assertFalse(dataframe["Code"].isna().any())
        self.assertTrue(dataframe["Code"].is_unique)

    def test_replaces_missing_duplicate_and_existing_codes_with_sequence(self):
        result = self.transform()
        self.assert_codes(result["dataframe"])
        self.assertEqual(result["dataframe"]["Name"].tolist(), self.source["Clinic Name"].tolist())

    def test_numbers_only_retained_rows_and_preserves_source_termination_matching(self):
        source = self.source.copy()
        source.loc[0, "Clinic Name"] = None
        source["Status"] = [None, None, "TERMINATED", None]
        source["Postal Code"] = "123456"
        result = self.transform(source, {("P100-1", "123456")})
        self.assert_codes(result["dataframe"])
        self.assertEqual(result["dataframe"]["Name"].tolist(), ["Beta Clinic"])
        self.assertEqual(result["terminated_clinics_filtered"], 2)
        self.assertIn("P100-1", result["filtered_provider_codes"])

    def test_missing_code_column_preserves_alignment_after_empty_row_filter(self):
        source = self.source.drop(columns=["Clinic ID", "S/N"])
        source.loc[0, "Clinic Name"] = None
        result = self.transform(source)
        self.assert_codes(result["dataframe"])
        self.assertEqual(result["dataframe"]["Name"].tolist(), [
            "Beta Clinic", "Gamma Clinic", "Delta Clinic",
        ])

    def test_all_listing_types_export_sequential_numeric_codes(self):
        sheet_names = ["GP Panel", "Dental Panel", "TCM Panel", "SP Clinic", "Provider List"]
        with pd.ExcelWriter(self.input_path, engine="openpyxl") as writer:
            for sheet in sheet_names:
                self.source.to_excel(writer, sheet_name=sheet, index=False)
        result = ExcelTransformer.transform_excel_multi_sheet(
            self.input_path, self.temp_dir.name, "test-codes", use_google_api=False,
        )
        self.assertTrue(result["success"], result.get("message"))
        self.assertEqual(len(result["output_files"]), len(sheet_names))
        for filename in result["output_files"]:
            with self.subTest(filename=filename):
                workbook = load_workbook(os.path.join(self.temp_dir.name, filename))
                self.addCleanup(workbook.close)
                rows = list(workbook.active.values)
                self.assertEqual(rows[0][0], "Code")
                self.assertEqual([row[0] for row in rows[1:]], [1, 2, 3, 4])
                self.assertTrue(all(workbook.active.cell(row, 1).data_type == "n" for row in range(2, 6)))

    def test_country_split_restarts_sequence_in_each_export(self):
        source = self.source.copy()
        source.loc[[1, 3], "Address"] = "2 JALAN TEST 80000 JOHOR MALAYSIA"
        source.loc[[1, 3], "Postal Code"] = "80000"
        source.to_excel(self.input_path, sheet_name="GP Panel", index=False)
        result = ExcelTransformer.transform_excel_multi_sheet(
            self.input_path, self.temp_dir.name, "test-split", use_google_api=False,
        )
        self.assertTrue(result["success"], result.get("message"))
        self.assertEqual(len(result["output_files"]), 2)
        for filename in result["output_files"]:
            with self.subTest(filename=filename):
                workbook = load_workbook(os.path.join(self.temp_dir.name, filename))
                self.addCleanup(workbook.close)
                rows = list(workbook.active.values)
                self.assertEqual([row[0] for row in rows[1:]], [1, 2])
                expected_names = ["Alpha Clinic", "Gamma Clinic"] if "Singapore" in filename else ["Beta Clinic", "Delta Clinic"]
                self.assertEqual([row[1] for row in rows[1:]], expected_names)

    def test_export_renumbers_filtered_indices_without_mutating_input(self):
        dataframe = pd.DataFrame({
            "Code": [7, None, 7],
            "Name": ["Alpha Clinic", "Beta Clinic", "Gamma Clinic"],
            "PostalCode": ["012345"] * 3,
        }, index=[2, 5, 8])
        original = dataframe.copy()
        output_path = os.path.join(self.temp_dir.name, "export.xlsx")
        ExcelTransformer.write_excel_with_text_postal_codes(dataframe, output_path)
        pd.testing.assert_frame_equal(dataframe, original)
        workbook = load_workbook(output_path)
        self.addCleanup(workbook.close)
        rows = list(workbook.active.values)
        self.assertEqual([row[0] for row in rows[1:]], [1, 2, 3])
        self.assertEqual([row[2] for row in rows[1:]], ["012345"] * 3)
        self.assertTrue(all(workbook.active.cell(row, 3).number_format == "@" for row in range(2, 5)))


if __name__ == "__main__":
    unittest.main()
